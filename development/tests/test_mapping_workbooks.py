"""Multi-tab attribute-mapping workbooks: cross-tab queries, status words in name columns,
conflicting duplicate mappings, and narrative tabs with embedded tables. Synthetic data only."""
from datetime import datetime

import openpyxl
import pytest

from backend.preprocessing.documents import ingest_document
from backend.preprocessing.profiling import column_quality, variant_groups
from backend.retrieval.search import search_documents
from backend.tools.documents import execute_document_tool

HEADER = ['Legacy Name', 'New Name', 'Cloud Name', 'Definition', 'Notes']


def mapping_rows(prefix, count):
    rows = [[f'{prefix} - Field {i}', f'{prefix} Field {i}', f'{prefix} Field {i}', f'Definition {i}', None]
            for i in range(count)]
    for i in range(0, count, 7):
        rows[i][1:3] = ['Out of Scope', 'Out of Scope']
        rows[i][4] = 'Not populated in the new system.'
    return rows


@pytest.fixture(scope='module')
def workbook(tmp_path_factory):
    wb = openpyxl.Workbook()
    history = wb.active
    history.title = 'Version History'
    history['A2'] = 'Mapping Workbook Title'
    history['A4'] = 'Revision History'
    history.append(['Date', 'Version', 'Description', 'Author'])
    history.append([datetime(2026, 1, 5), 3, 'Added Billing tab', 'Dana Park'])
    history.append([datetime(2026, 3, 9), 4, 'Removed retired fields', 'Lee Ortiz'])
    assumptions = wb.create_sheet('Assumptions')
    assumptions.append(['Pending means availability is still being confirmed.'])
    for name, prefix, count in [('Billing', 'Invoice', 40), ('Shipping', 'Parcel', 45)]:
        sheet = wb.create_sheet(name)
        sheet.append(HEADER)
        for row in mapping_rows(prefix, count):
            sheet.append(row)
    shipping = wb['Shipping']
    shipping.append(['Invoice - Field 3', 'Invoice Field 3', 'Invoice Field 3', 'Definition 3', None])
    shipping.append(['Parcel - Field 2', 'Pending', 'Pending', 'Definition 2', 'Availability being confirmed.'])
    areas = wb.create_sheet('Subject Areas')
    areas.append(['Area', 'Cloud', 'Wave'])
    areas.append(['Billing', 'Available', 'Wave 1'])
    areas.append(['Shipping', 'Partial', 'Wave 2'])
    areas.append(['Returns', 'Partial', 'Wave 2'])
    path = tmp_path_factory.mktemp('mapping') / 'mapping.xlsx'
    wb.save(path)
    return ingest_document(path.name, path.read_bytes())


def query(doc, **params):
    return execute_document_tool('table_tool', {'document_id': doc.id, 'name': 'query_table', 'parameters': params},
                                 {doc.id: doc})['table_result']


def test_query_all_tabs_with_provenance(workbook):
    result = query(workbook, sheet_names=['*'], filters=[{'column': 'Legacy Name', 'op': 'eq',
                   'value': 'Invoice - Field 3'}], select=['New Name'])
    assert [(r['_sheet'], r['_source_row'], r['New Name']) for r in result['rows']] == [
        ('Billing', 5, 'Invoice Field 3'), ('Shipping', 47, 'Invoice Field 3')]
    assert 'Subject Areas' in result['skipped_sheets']
    per_tab = query(workbook, sheet_names=['*'], filters=[{'column': 'New Name', 'op': 'eq', 'value': 'Out of Scope'}],
                    group_by=['_sheet'], aggregations=[{'op': 'count_rows', 'as': 'n'}])
    assert {r['_sheet']: r['n'] for r in per_tab['rows']} == {'Billing': 6, 'Shipping': 7}
    absent = query(workbook, sheet_names=['*'], filters=[{'column': 'Legacy Name', 'op': 'eq', 'value': 'Nope'}],
                   aggregations=[{'op': 'count_rows', 'as': 'n'}])
    assert absent['rows'] == [{'n': 0}] and absent['sheets_queried'] == ['Billing', 'Shipping']
    assert 'error' in query(workbook, sheet_names=['Missing'], select=['New Name'])


def test_single_tab_miss_points_to_other_tabs(workbook):
    result = query(workbook, sheet_name='Billing', filters=[{'column': 'Legacy Name', 'op': 'eq',
                   'value': 'Parcel - Field 1'}])
    assert result['matched_rows'] == 0
    assert any("['Shipping']" in note and 'sheet_names' in note for note in result['notes'])


def test_value_in_another_column_is_reported(workbook):
    result = query(workbook, sheet_name='Billing', filters=[{'column': 'Notes', 'op': 'contains', 'value': 'out of scope'}])
    assert result['matched_rows'] == 0
    assert result['filter_diagnostics'][0]['value_found_in_columns'] == {'New Name': 6, 'Cloud Name': 6}


def test_conflicting_mappings_flagged_only_for_identifier_keys(workbook):
    conflict = query(workbook, sheet_names=['*'], filters=[{'column': 'Legacy Name', 'op': 'eq',
                     'value': 'Parcel - Field 2'}], select=['New Name'])
    assert any("disagree in ['New Name']" in note for note in conflict['notes'])
    category = query(workbook, sheet_name='Subject Areas', filters=[{'column': 'Cloud', 'op': 'eq', 'value': 'Partial'}],
                     select=['Area'])
    assert category['matched_rows'] == 2 and category['notes'] == []


def test_status_words_in_name_columns_are_hinted(workbook):
    hints = workbook.summary()['sheets']['Shipping']['column_hints']
    assert hints['New Name']['frequent_values'] == {'Out of Scope': 7}
    assert 'possible_variant_groups' not in hints.get('Legacy Name', {})


@pytest.mark.parametrize('values,grouped', [
    (['Account - Date', 'Account - End Date'], False),
    (['Pack - Amount', 'Package - Amount'], False),
    (['Order Date', 'Order Date Display'], False),
    (['Customer Acount', 'Customer Account'], True),
    (['fail', 'failed'], True),
])
def test_added_words_and_extended_terms_are_distinct(values, grouped):
    assert bool(variant_groups({v: 3 for v in values})) is grouped


def test_mostly_unique_name_columns_skip_variant_checks():
    import pandas as pd
    names = pd.Series([f'Item {i} - Date' for i in range(60)] + [f'Item {i} - End Date' for i in range(60)],
                      dtype='string')
    assert 'possible_variant_groups' not in column_quality(names)


def test_narrative_tab_tables_read_as_rows(workbook):
    lines = [b['text'] for b in workbook.blocks if b['location']['sheet'] == 'Version History']
    assert lines == ['Mapping Workbook Title', 'Revision History', 'Date | Version | Description | Author',
                     'Date: 2026-01-05; Version: 3; Description: Added Billing tab; Author: Dana Park',
                     'Date: 2026-03-09; Version: 4; Description: Removed retired fields; Author: Lee Ortiz']
    hit = search_documents({workbook.id: workbook}, 'version 4 author')['matches'][0]
    assert 'Lee Ortiz' in hit['text'] and hit['location'] == {'sheet': 'Version History', 'row': 7}
    page = execute_document_tool('table_tool', {'document_id': workbook.id, 'name': 'read_sheet', 'parameters': {
        'sheet_name': 'Version History', 'offset': 3, 'limit': 1}}, {workbook.id: workbook})['table_result']
    assert page['rows'] == [{'source_row': 6, 'text': lines[3]}] and page['total_rows'] == 5


def test_query_narrative_tab_uses_row_text(workbook):
    result = query(workbook, sheet_name='Version History',
                   filters=[{'column': 'text', 'op': 'contains', 'value': 'Version: 4'}])
    assert result['rows'] == [{'source_row': 7, 'text': 'Date: 2026-03-09; Version: 4; Description: Removed '
                                                          'retired fields; Author: Lee Ortiz'}]
    assert result['sources'][0]['source_rows'] == [7]


def test_missing_column_error_suggests_sheet_column(workbook):
    result = query(workbook, sheet_names=['Billing', 'Shipping'], group_by=['Subject Area'],
                   aggregations=[{'op': 'count_rows', 'as': 'n'}])
    assert '_sheet' in result['error'] and 'Legacy Name' in result['error']
