from io import BytesIO
from pathlib import Path

import pytest

from backend.preprocessing.documents import ingest_document
from backend.retrieval.search import search_documents
from backend.tools.documents import execute_document_tool
from backend.agents.document_agent import run_document_agent

FIXTURES = Path(__file__).resolve().parents[1] / 'fixtures'


def make_docs():
    docs = [ingest_document(p.name,p.read_bytes()) for p in sorted((FIXTURES/'documents').glob('*.txt'))]
    return {d.id:d for d in docs}


def test_file_identity_includes_content_and_name():
    assert ingest_document('a.txt',b'one').id != ingest_document('a.txt',b'two').id
    assert ingest_document('a.txt',b'one').id != ingest_document('b.txt',b'one').id


def test_selected_scope_cannot_be_overridden_by_tool():
    docs = make_docs()
    first = next(iter(docs))
    other = next(key for key in docs if key != first)
    subset = {first:docs[first]}
    with pytest.raises(ValueError,match='selected'):
        execute_document_tool('read_document',{'document_id':other},subset)
    with pytest.raises(ValueError,match='unavailable'):
        execute_document_tool('search_documents',{'query':'rates','document_ids':[other]},subset)


def test_lexical_retrieval_has_source_locations():
    docs = make_docs()
    result = search_documents(docs,'USD 88.00')
    assert result['matches'][0]['filename'] == 'bingle_dingle_amendment.txt'
    assert 'paragraph' in result['matches'][0]['location']
    assert '88.00' in result['matches'][0]['text']


def test_no_generated_code_tool_in_document_app(insurance_workbook):
    path = insurance_workbook.with_name('insurance_claims.csv')
    doc = ingest_document(path.name,path.read_bytes())
    with pytest.raises(ValueError,match='disabled'):
        execute_document_tool('table_tool',{'document_id':doc.id,'name':'analyze_data'}, {doc.id:doc})


def test_document_app_can_query_selected_table(insurance_workbook):
    path = insurance_workbook.with_name('insurance_claims.csv')
    doc = ingest_document(path.name,path.read_bytes())
    result = execute_document_tool('table_tool', {'document_id':doc.id,'name':'query_table',
        'parameters':{'sheet_name':'Sheet1','filters':[{'column':'Claim ID','op':'eq','value':'0012'}]}}, {doc.id:doc})
    assert result['table_result']['rows'][0]['Claim ID'] == '0012'


def test_read_docx_body_and_tables():
    from docx import Document
    file = Document()
    file.add_heading('Service rules',level=1)
    file.add_paragraph('Session rate is USD 88.00.')
    table = file.add_table(rows=1,cols=2)
    table.cell(0,0).text='Service'
    table.cell(0,1).text='Price'
    buffer = BytesIO()
    file.save(buffer)
    doc = ingest_document('rules.docx',buffer.getvalue())
    assert any('88.00' in b['text'] for b in doc.blocks)
    assert any('table' in b['location'] for b in doc.blocks)
    assert all('page' not in b['location'] for b in doc.blocks)


def test_scanned_pdf_reports_missing_text():
    from pypdf import PdfWriter
    pdf = PdfWriter()
    pdf.add_blank_page(width=612,height=792)
    buffer = BytesIO()
    pdf.write(buffer)
    doc = ingest_document('scan.pdf',buffer.getvalue())
    assert not doc.blocks
    assert any('OCR' in w for w in doc.warnings)


def test_text_pdf_page_citations():
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    pdf = PdfWriter()
    page = pdf.add_blank_page(width=612,height=792)
    font = DictionaryObject({NameObject('/Type'):NameObject('/Font'), NameObject('/Subtype'):NameObject('/Type1'),
                             NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 50 700 Td (Session rate USD 88.00) Tj ET')
    page[NameObject('/Contents')] = stream
    buffer=BytesIO()
    pdf.write(buffer)
    doc=ingest_document('rates.pdf',buffer.getvalue())
    assert doc.blocks[0]['location'] == {'page':1}
    assert '88.00' in doc.blocks[0]['text']


def test_agent_sees_only_selected_documents(native_script):
    import json
    docs=make_docs()
    selected=next(key for key,d in docs.items() if d.name=='bingle_dingle_amendment.txt')
    requests=native_script([{'tool':'read_document','parameters':{'document_id':selected}},
                           {'tool':'answer','parameters':{}},'USD 88.00 per session [E1].'])
    result=run_document_agent(docs,[selected],'What is the session rate?')
    assert result['answer']=='USD 88.00 per session [E1].'
    payload=json.loads(requests[0]['messages'][0]['content'][0]['text'])
    assert all(d['document_id']==selected for d in payload['selected_documents'])
    assert len(result['evidence'])==1


def test_invalid_citation_is_flagged(native_script):
    docs=make_docs()
    selected=next(iter(docs))
    native_script([{'tool':'read_document','parameters':{'document_id':selected}},
                   {'tool':'answer','parameters':{}},'A statement [E99].'])
    result=run_document_agent(docs,[selected],'Question')
    assert any('unrecognized' in w for w in result['limitations'])
