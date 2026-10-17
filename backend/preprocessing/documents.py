"""Document ingestion and optional persistence for selected document Q&A."""
from backend.shared.models import Document
from backend.shared.rows import narrative_rows, row_location, row_text
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_PAGES = 200
MAX_BLOCKS = 2000
SUPPORTED = {'.pdf', '.docx', '.xlsx', '.xls', '.csv', '.txt', '.md'}


def _add_text(doc, text, location):
    text = text.replace('\x00', '').strip()
    for offset in range(0, len(text), 1400):
        if len(doc.blocks) >= MAX_BLOCKS:
            warning = f'Extraction limited to {MAX_BLOCKS} text blocks; some content is unavailable.'
            if warning not in doc.warnings:
                doc.warnings.append(warning)
            return
        # Overlap allows matches spanning a chunk boundary, without crossing source pages.
        doc.blocks.append({'block_id':f'B{len(doc.blocks)+1}', 'text':text[offset:offset+1600],
                           'location':location})


def ingest_document(name: str, content: bytes, *, store=None, enrich=None, hint=None, embed=None,
                    table_options=None) -> Document:
    """Extract a document's searchable text (and, for spreadsheets, per-sheet tables), then optionally profile it,
    add search hints, embed its blocks and save it.

    hint           e.g. search_hints.make_hinter(): per-block search hints for keyword search

    embed          e.g. embeddings.make_embedder(): per-block vectors for semantic search

    table_options  spreadsheet preprocessing options: encoding, delimiter (CSV) and sheet overrides

    Runs in its own trace (ingest_document), archived when it ends; a failure is recorded on it and re-raised.
    """
    from backend.observability import tracing
    from backend.observability.archive import configured_archive, document_summary
    with tracing.span('workflow', 'ingest_document', archive=configured_archive()) as current:
        tracing.annotate(current, metadata={
            'filename': name, 'kind': Path(name).suffix.lower(), 'bytes': len(content or b''),
            'content_sha256': sha256(content or b'').hexdigest(), 'profile_requested': enrich is not None,
            'search_hints_requested': hint is not None, 'embeddings_requested': embed is not None,
            'saved': store is not None, 'group_name': getattr(store, 'group', None)})
        doc = _ingest(name, content, store, enrich, hint, embed, table_options)
        profile = doc.profile or {}
        tracing.annotate(current, metadata={'document': document_summary(doc),
                                            'profile_generated': bool(profile),
                                            'profile_model_id': profile.get('meta', {}).get('model_id')})
        return doc


def _ingest(name, content, store, enrich, hint, embed, table_options):
    name = name.replace('\\','/').rsplit('/',1)[-1]
    extension = Path(name).suffix.lower()
    if extension not in SUPPORTED:
        raise ValueError(f'Unsupported file type: {extension}')
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError('File must be nonempty and at most 20 MB')
    doc = Document(sha256(name.encode() + b'\0' + content).hexdigest()[:16], name, extension[1:])
    if extension == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted and not reader.decrypt(''):
            raise ValueError('Password-protected PDFs must be unlocked before upload')
        if len(reader.pages) > MAX_PAGES:
            doc.warnings.append(f'Only the first {MAX_PAGES} PDF pages were extracted.')
        for index, page in enumerate(reader.pages[:MAX_PAGES], 1):
            text = page.extract_text() or ''
            if not text.strip():
                doc.warnings.append(f'Page {index} has no extractable text; scanned pages require OCR.')
            else:
                _add_text(doc, text, {'page':index})
        doc.warnings.append('PDF text order and table layout may differ from the original; verify cited pages.')
    elif extension == '.docx':
        from docx import Document as WordDocument
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        word = WordDocument(BytesIO(content))
        paragraph_number, table_number, section = 0, 0, ''
        for element in word.element.body.iterchildren():
            if element.tag.endswith('}p'):
                paragraph = Paragraph(element, word)
                paragraph_number += 1
                if paragraph.style and paragraph.style.name.startswith('Heading'):
                    section = paragraph.text
                _add_text(doc, paragraph.text, {'paragraph':paragraph_number, 'section':section})
            elif element.tag.endswith('}tbl'):
                table_number += 1
                table = Table(element, word)
                for row_number, row in enumerate(table.rows, 1):
                    _add_text(doc, ' | '.join(cell.text for cell in row.cells),
                              {'table':table_number, 'row':row_number, 'section':section})
        doc.warnings.append('Word citations use paragraphs/tables, not page numbers. Headers, footers, text boxes and tracked changes are not fully extracted.')
    elif extension in {'.xlsx', '.xls', '.csv'}:
        from backend.preprocessing.tables import preprocess_table
        from backend.storage.tables import MemoryTableFolder
        files, doc.table_metadata = preprocess_table(name, content, **(table_options or {}))
        doc.table_folder = MemoryTableFolder(files)
        index_tables(doc)
    else:
        text = content.decode('utf-8-sig')
        for number, paragraph in enumerate(re.split(r'\n\s*\n', text), 1):
            _add_text(doc, paragraph, {'paragraph':number})
    if not doc.blocks and not doc.table_folder:
        doc.warnings.append('No searchable text was extracted from this file.')
    if enrich is not None:
        # Optional LLM semantic profile (backend/agents/enrichment.py). Sends the extracted text to the model.
        try:
            doc.profile = enrich(doc)
        except Exception as exc:
            doc.warnings.append(f'Semantic profile not generated: {type(exc).__name__}: {exc}')
    if hint is not None and doc.blocks:
        # Optional LLM search hints (backend/agents/search_hints.py). Sends the extracted text to the model.
        try:
            doc.search_hints = hint(doc)
        except Exception as exc:
            doc.warnings.append(f'Search hints not generated: {type(exc).__name__}: {exc}')
    if embed is not None and doc.blocks:
        # Optional passage embeddings (backend/agents/embeddings.py). Sends the extracted text to the model.
        try:
            doc.embeddings = embed(doc)
        except Exception as exc:
            doc.warnings.append(f'Embeddings not generated: {type(exc).__name__}: {exc}')
    if store is not None:
        store.save(doc, content)
    return doc


def index_tables(doc):
    """Text blocks for a spreadsheet: narrative sheets and the notes around tables.

    Data-sheet rows are not copied into blocks: the search tool and the table tools read every row from the
    sheet's CSV, so a copy would only crowd the block limit.
    """
    from backend.storage.tables import load_table
    for sheet, meta in doc.table_metadata['sheets'].items():
        if meta.get('sheet_type') == 'metadata':
            frame = load_table(doc.table_folder, meta)
            if {'source_row', 'source_column', 'text'} <= set(frame.columns):
                # Narrative sheets are indexed in full, one block per sheet row (not per cell).
                for number, text in narrative_rows(frame.rows):
                    _add_text(doc, text, {'sheet': sheet, 'row': number})
            else:
                for i, row in enumerate(frame.rows):
                    _add_text(doc, row_text(row), row_location(sheet, meta, i, row))
        for row in meta.get('context_rows', []):
            _add_text(doc, ' | '.join(row['values']), {'sheet':sheet, 'row':row['source_row']})
    doc.warnings.append('Spreadsheet ingestion uses heuristic headers; inspect ambiguous layouts. Formulas need cached values. Use table tools, not text snippets, for counts and totals.')
